"""
import_xlsx.py
==============
One-time import of the old Google Sheet (downloaded as FantasyB1G.xlsx) into
the site's JSON data files. You shouldn't need to run this again unless you
want to re-import history from the spreadsheet.

    pip install openpyxl
    python tools/import_xlsx.py path/to/FantasyB1G.xlsx

Writes:
    data/seasons/<year>.json   completed regular seasons (2016..2026)
    data/mm/<year>.json        completed March Madness contests
    data/extras.json           HOF, ballots, schools
    league/<year>.json         the in-progress season's config (draft so far)
"""
import json
import os
import re
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import MANAGERS, LAST_TO_FIRST, season_label, write_json  # noqa: E402

FINAL_SEASONS = range(2016, 2027)
CURRENT_SEASON = 2027
FINAL_MM = [y for y in range(2016, 2027)]
# sheet abbreviations -> the ESPN/CBBD school names used everywhere else
SCHOOL_NAMES = {"Michigan St": "Michigan State", "Ohio St": "Ohio State"}


def clean_name(raw, flag=None):
    """Strip the '*' bench marker (prefix, suffix, or in the flag column)."""
    name = re.sub(r"\s+", " ", str(raw)).strip()
    name = re.sub(r"\s*\(\d+\)$", "", name)  # March Madness sheets append the seed: "Coby White (1)"
    bench = bool(flag and "*" in str(flag))
    if name.startswith("*"):
        bench, name = True, name[1:].strip()
    if name.endswith("*"):
        bench, name = True, name[:-1].strip()
    return name, bench


def years_text(v):
    """Excel stores a lone year as a float (2023.0)."""
    return str(int(v)) if isinstance(v, float) else str(v).strip()


def stat(v):
    if v is None or (isinstance(v, str) and v.strip() in ("", "-")):
        return None
    return int(round(float(v)))


def parse_blocks(rows, managers, max_row=None):
    """Return {manager: [[row, ...] per game slot]} from the side-by-side
    'Game # | flag | Players | Pts Rebs Asts Stls Blks | Opponent' blocks."""
    header = rows[0]
    starts = [i for i, v in enumerate(header) if v == "Game #"]
    out = {}
    for mgr, c0 in zip(managers, starts):
        games = {}
        for r in rows[1:max_row]:
            if len(r) <= c0 + 8:
                continue
            g, flag, nm = r[c0], r[c0 + 1], r[c0 + 2]
            if not isinstance(g, (int, float)) or nm in (None, "") or str(nm).startswith("sub"):
                continue
            name, bench = clean_name(nm, flag)
            s = [stat(r[c0 + 3 + i]) for i in range(5)]
            opp = r[c0 + 8]
            row = {"p": name, "o": str(opp).strip() if opp else None}
            if any(x is not None for x in s):
                row["s"] = [x or 0 for x in s]
            else:
                row["s"] = None
            if bench:
                row["b"] = 1
            games.setdefault(int(g), []).append(row)
        out[mgr] = [games[g] for g in sorted(games)]
    return out


def roster_segments(games):
    """Per draft slot, the run of players who held it: [{p, from, to}]."""
    nslots = max(len(g) for g in games) if games else 0
    segs = []
    for slot in range(nslots):
        cur = []
        for gi, g in enumerate(games, start=1):
            if slot >= len(g):
                continue
            p = g[slot]["p"]
            if p is None or g[slot].get("fill"):
                continue  # 2021 COVID fill-ins / empty rows aren't roster moves
            if cur and cur[-1]["p"] == p:
                cur[-1]["to"] = gi
            else:
                cur.append({"p": p, "from": gi, "to": gi})
        segs.append(cur)
    return segs


def side_notes(rows):
    """Free-text notes typed under the week records (draft order, buy-ins...)."""
    notes = []
    for r in range(85, min(96, len(rows))):
        line = " ".join(str(v).strip() for v in rows[r][1:5] if v not in (None, ""))
        if line and "$" not in line:  # buy-ins / payouts aren't tracked on the site
            notes.append(line.replace("Orer", "Order"))
    return notes


def draft_order(notes):
    text = " ".join(notes)
    m = re.search(r"\b([MWBF]{4})\b", text) or re.search(r"\b([MWBF]),([MWBF]),([MWBF]),([MWBF])\b", text)
    if m:
        return "".join(m.groups()) if len(m.groups()) > 1 else m.group(1)
    first = {"Falan": "F", "Brett": "B", "Waldo": "W", "Michael": "M"}
    letters = [first[n] for n in notes if n in first]
    return "".join(letters) if len(letters) == 4 else None


def import_season(wb, year):
    rows = [[c.value for c in r] for r in wb[str(year)].iter_rows()]
    managers = [LAST_TO_FIRST[rows[0][c].strip()] for c in range(1, 5)]
    teams_games = parse_blocks(rows, managers, max_row=280)
    notes = side_notes(rows)
    covid = {}
    if str(rows[29][0]).strip().upper() == "COVID":
        # 2021: each team got a fill-in player whose games covered rostered
        # players whose seasons ended early. The sheet logs them as e.g.
        # "AL DURHAM 4" (games counting down) -- fold back to the real name.
        covid = {managers[i]: str(rows[29][1 + i]).strip() for i in range(4)}
        notes.insert(0, "COVID season (27 games). Late in the year each team added a fill-in player whose games "
                        "covered rostered players whose seasons ended early: "
                        + ", ".join(f"{p} ({m})" for m, p in covid.items()) + ".")
        for mgr, games in teams_games.items():
            for g in games:
                for row in g:
                    if re.fullmatch(r"[A-Z' .\-]+ \d", row["p"]):
                        row["p"], row["fill"] = covid[mgr], 1
                    elif row["p"].lower() == "benched":
                        row["p"] = None
    order = draft_order(notes)
    pickup_rows = []
    for r in range(65, 68):
        d = rows[r][0]
        if d is None:
            continue
        pickup_rows.append({
            "date": d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else str(d),
            "adds": {managers[i]: (str(rows[r][1 + i]).strip() if rows[r][1 + i] else None) for i in range(4)},
            "order": rows[r][5],
        })

    teams = {}
    for mgr in managers:
        games = teams_games[mgr]
        segs = roster_segments(games)
        draft = [s[0]["p"] for s in segs]
        pickups = []
        for slot, s in enumerate(segs):
            for prev, nxt in zip(s, s[1:]):
                pickups.append({"slot": slot + 1, "add": nxt["p"], "drop": prev["p"], "fromGame": nxt["from"]})
        # attach the sheet's pickup dates where the add matches
        for pk in pickups:
            for pr in pickup_rows:
                if (pr["adds"].get(mgr) or "").lower() == pk["add"].lower():
                    pk["date"] = pr["date"]
        teams[mgr] = {"draft": draft, "pickups": sorted(pickups, key=lambda x: x["fromGame"]), "games": games}

    return {
        "year": year,
        "label": season_label(year),
        "kind": "season",
        "status": "final",
        "managers": managers,
        "slots": max(len(t["games"]) for t in teams.values()),
        "draftOrder": order,
        "notes": notes,
        "teams": teams,
    }


def import_mm(wb, year):
    rows = [[c.value for c in r] for r in wb[f"{year}MM"].iter_rows()]
    managers = [str(rows[0][c]).strip() for c in range(1, 5)]
    teams_games = parse_blocks(rows, managers)
    teams = {}
    for mgr in managers:
        col = 1 + managers.index(mgr)
        draft = [clean_name(rows[r][col])[0] for r in range(1, 8) if rows[r][col]]
        games = teams_games.get(mgr, [])
        teams[mgr] = {"draft": draft, "games": games}
    cancelled = not any(t["games"] and any(r["s"] for g in t["games"] for r in g) for t in teams.values())
    return {
        "year": year,
        "label": f"{year} March Madness",
        "kind": "mm",
        "status": "cancelled" if cancelled else "final",
        "note": "Cancelled (COVID-19)" if cancelled else None,
        "managers": managers,
        "slots": 6,
        "teams": teams,
    }


def import_current_config(wb, year):
    rows = [[c.value for c in r] for r in wb[str(year)].iter_rows()]
    managers = [LAST_TO_FIRST[rows[0][c].strip()] for c in range(1, 5)]
    teams = {}
    for i, mgr in enumerate(managers):
        draft = [clean_name(rows[r][1 + i])[0] if rows[r][1 + i] else "" for r in range(1, 8)]
        teams[mgr] = {"draft": draft, "bench": [], "pickups": []}
    return {
        "season": year,
        "slots": 31,
        "draftOrder": draft_order(side_notes(rows)),
        "pickupDates": [rows[r][0].strftime("%Y-%m-%d") for r in range(65, 68) if rows[r][0]],
        "teams": teams,
        "players": {},
        "overrides": [],
    }


def import_extras(wb):
    # Player Index: school + displayed years for every player who's been rostered
    pi = [[c.value for c in r] for r in wb["Player Index"].iter_rows()]
    schools = {}
    for r in pi[1:]:
        if r[0] and r[1] and r[2] and str(r[0]) not in ("Hall of Fame", "Active"):
            school = str(r[2]).strip()
            schools[f"{str(r[0]).strip()} {str(r[1]).strip()}"] = SCHOOL_NAMES.get(school, school)
    # rostered but missing from the Player Index
    schools.setdefault("Isaiah Todd", "Michigan")
    schools.setdefault("Tarris Reed", "Michigan")

    hof_rows = [[c.value for c in r] for r in wb["Hall of Fame"].iter_rows()]
    members = []
    for r in hof_rows[1:]:
        if r[0] and r[1] and isinstance(r[3], str) and "(" in r[3]:
            yr, pct = re.match(r"(\d{4})\s*\((\d+)%\)", r[3]).groups()
            members.append({"player": r[0].strip(), "owner": r[1].strip(), "years": years_text(r[2]),
                            "inducted": int(yr), "vote": int(pct)})
    ballots = {}
    header = hof_rows[0]
    for c, v in enumerate(header):
        mt = re.match(r"(\d{4}) Nominations", str(v or ""))
        if not mt:
            continue
        noms = []
        for r in hof_rows[1:26]:
            if c < len(r) and r[c]:
                noms.append({"player": str(r[c]).strip(), "owner": str(r[c + 1]).strip(),
                             "years": years_text(r[c + 2]), "yearsOnBallot": r[c + 3]})
        ballots[mt.group(1)] = noms

    return {
        "schools": schools,
        "hof": {
            "members": members,
            "ballots": ballots,
            "rules": [
                "Eligible after a player's fantasy career ends: more than 20 games played and more than 500 total points.",
                "3/4 of the vote gets a player in.",
                "3/4 of the vote takes a player off the ballot.",
                "A player can stay on the ballot for at most 4 years.",
            ],
        },
    }


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: python tools/import_xlsx.py FantasyB1G.xlsx")
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    for y in FINAL_SEASONS:
        write_json(os.path.join(ROOT, "data", "seasons", f"{y}.json"), import_season(wb, y))
        print("season", y)
    for y in FINAL_MM:
        write_json(os.path.join(ROOT, "data", "mm", f"{y}.json"), import_mm(wb, y))
        print("mm", y)
    write_json(os.path.join(ROOT, "data", "extras.json"), import_extras(wb))
    cfg_path = os.path.join(ROOT, "league", f"{CURRENT_SEASON}.json")
    if os.path.exists(cfg_path):
        print("kept existing", cfg_path)
    else:
        write_json(cfg_path, import_current_config(wb, CURRENT_SEASON), pretty=True)
        print("wrote", cfg_path)


if __name__ == "__main__":
    main()
